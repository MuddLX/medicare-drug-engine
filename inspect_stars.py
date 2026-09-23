# inspect_stars.py — dumps the structure of the CMS Summary Ratings XLSX so we can
# map its columns before building the parser.
# Run from repo root:  python inspect_stars.py     (needs: pip install openpyxl)
import glob
import os

import openpyxl

folder = os.path.join("Star Ratings", "Master Tables Individual Files 8 17 26")
matches = glob.glob(os.path.join(folder, "*Summary Ratings*.xlsx"))
if not matches:
    print("Summary Ratings xlsx not found under", folder)
    raise SystemExit

path = matches[0]
print("File:", os.path.basename(path))
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("Sheets:", wb.sheetnames)

ws = wb[wb.sheetnames[0]]
print(f"\nFirst 15 rows of '{ws.title}' (trailing blanks trimmed):")
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
    vals = list(row)
    while vals and vals[-1] is None:
        vals.pop()
    print(f"  row {i}: {vals}")
